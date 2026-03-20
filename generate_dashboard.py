"""
generate_dashboard.py
=====================
Reads from the two database tables (bess_revenues, bess_soc_model) and
produces BESS_Dashboard.xlsx — a single-sheet executive summary designed
for non-technical readers. No terminal output required to understand results.

Run after data_retrieval_fingrid.py and data_analytics.py:
    python generate_dashboard.py
"""

import os
import math
import pandas as pd
import sqlalchemy
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv

load_dotenv()
DB_TYPE = os.getenv("DB_TYPE", "sqlite").lower()

BESS_MW  = 30
BESS_MWH = 36
CAPEX    = 12_000_000
WACC     = 0.08
OPS_DISC = 0.20

DASHBOARD_FILE = "BESS_Dashboard.xlsx"

# ---------------------------------------------------------------------------
# Colour palette
# ---------------------------------------------------------------------------
C = {
    "navy":        "1F4E79",
    "mid_blue":    "2E75B6",
    "light_blue":  "D6E4F0",
    "green":       "375623",
    "light_green": "E2EFDA",
    "amber":       "7F6000",
    "light_amber": "FFF2CC",
    "red":         "842019",
    "light_red":   "FCE4D6",
    "white":       "FFFFFF",
    "light_grey":  "F5F5F5",
    "mid_grey":    "D9D9D9",
    "dark_grey":   "595959",
    "teal":        "2E8B57",
}

# ---------------------------------------------------------------------------
# DB helpers
# ---------------------------------------------------------------------------
def get_engine(fallback_to_sqlite=True):
    """
    Returns a working DB engine. If DB_TYPE=postgres and the container is
    unreachable, automatically falls back to SQLite so downstream scripts
    never crash due to a stopped Docker container.
    """
    if DB_TYPE == "postgres":
        try:
            from sqlalchemy import text as sa_text
            url = (f"postgresql://{os.getenv('POSTGRES_USER')}:"
                   f"{os.getenv('POSTGRES_PASSWORD')}@"
                   f"{os.getenv('POSTGRES_HOST')}:{os.getenv('POSTGRES_PORT')}/"
                   f"{os.getenv('POSTGRES_DB')}")
            engine = sqlalchemy.create_engine(url)
            with engine.connect() as conn:
                conn.execute(sa_text("SELECT 1"))
            return engine
        except Exception as e:
            if fallback_to_sqlite:
                short = str(e).split("\n")[0][:100]
                print(f"  [!] PostgreSQL unavailable ({short})")
                print("      Falling back to SQLite — run: docker-compose up -d")
                return sqlalchemy.create_engine("sqlite:///bess_model.db")
            raise
    return sqlalchemy.create_engine("sqlite:///bess_model.db")


def load_data():
    engine = get_engine()
    df_p = pd.read_sql("SELECT * FROM bess_revenues", engine)
    df_p["Timestamp"] = pd.to_datetime(df_p["Timestamp"])
    df_p = df_p.sort_values("Timestamp").set_index("Timestamp")
    df_s = pd.read_sql("SELECT * FROM bess_soc_model", engine)
    df_s["Timestamp"] = pd.to_datetime(df_s["Timestamp"])
    df_s = df_s.sort_values("Timestamp").set_index("Timestamp")
    return df_p, df_s


# ---------------------------------------------------------------------------
# Financial calcs
# ---------------------------------------------------------------------------
def npv(rate, cfs):
    try:
        return sum(cf / (1 + rate) ** t for t, cf in enumerate(cfs))
    except (OverflowError, ZeroDivisionError):
        return float("inf")


def xirr(cfs, guess=0.05, tol=1e-7, max_iter=500):
    import numpy as np
    tests = np.linspace(-0.9, 2.0, 30)
    signs = set(math.copysign(1, npv(r, cfs)) for r in tests
                if not math.isnan(npv(r, cfs)))
    if len(signs) < 2:
        return float("nan")
    rate = max(-0.9999, min(guess, 10.0))
    for _ in range(max_iter):
        f  = npv(rate, cfs)
        d  = sum(-t * cf / (1 + rate) ** (t + 1) for t, cf in enumerate(cfs))
        if d == 0:
            return float("nan")
        r2 = max(-0.9999, min(rate - f / d, 10.0))
        if abs(r2 - rate) < tol:
            return r2
        rate = r2
    return float("nan")


def payback(cfs):
    cum = 0.0
    for t, cf in enumerate(cfs):
        prev, cum = cum, cum + cf
        if t > 0 and cum >= 0:
            return (t - 1) + abs(prev) / cf if cf else t
    return float("inf")


def build_cf(rev_y1, capex=CAPEX, opex=180_000, opex_var=0.005,
             repl_yr=12, repl_cost=3_500_000, life=20,
             tax=0.20, depr_yrs=10, growth=0.03, deg=0.02):
    rows, depr = [], capex / depr_yrs
    for yr in range(1, life + 1):
        rev    = rev_y1 * (1 - deg) ** (yr - 1) * (1 + growth) ** (yr - 1)
        o      = opex + opex_var * rev
        ebitda = rev - o
        d      = depr if yr <= depr_yrs else 0
        if yr > repl_yr:
            d += repl_cost / (life - repl_yr)
        ebit = ebitda - d
        t_   = max(0, ebit * tax)
        repl = repl_cost if yr == repl_yr else 0
        fcf  = (ebit - t_) + d - repl
        rows.append({"yr": yr, "rev": rev, "fcf": fcf, "cumfcf": 0})
    cum = 0
    for r in rows:
        cum += r["fcf"]
        r["cumfcf"] = cum - capex
    return rows


# ---------------------------------------------------------------------------
# Style helpers
# ---------------------------------------------------------------------------
def fill(hex_):
    return PatternFill("solid", fgColor=hex_)

def font(bold=False, size=10, color="000000", italic=False):
    return Font(bold=bold, size=size, color=color, italic=italic, name="Arial")

def align(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def border_thin():
    s = Side(style="thin", color="D0D0D0")
    return Border(left=s, right=s, top=s, bottom=s)

def border_none():
    n = Side(style=None)
    return Border(left=n, right=n, top=n, bottom=n)

def set_col_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def spacer_row(ws, row, height=10):
    ws.row_dimensions[row].height = height

def merge_cell(ws, rng, value="", bg=None, fg="000000",
               bold=False, size=10, h="left", v="center",
               italic=False, wrap=False, border=False):
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    c.value = value
    if bg:
        c.fill = fill(bg)
    c.font      = font(bold=bold, size=size, color=fg, italic=italic)
    c.alignment = align(h=h, v=v, wrap=wrap)
    c.border    = border_thin() if border else border_none()
    return c


# ---------------------------------------------------------------------------
# KPI card — 3 rows: label / big value / note
# ---------------------------------------------------------------------------
def kpi_card(ws, top_row, col, label, value, note,
             label_bg, value_bg, value_size=18):
    cl = get_column_letter(col)
    cr = get_column_letter(col + 1)
    merge_cell(ws, f"{cl}{top_row}:{cr}{top_row}",
               label, bg=label_bg, fg=C["white"],
               bold=False, size=8, h="center", italic=True)
    merge_cell(ws, f"{cl}{top_row+1}:{cr}{top_row+1}",
               value, bg=value_bg, fg=C["white"],
               bold=True, size=value_size, h="center")
    merge_cell(ws, f"{cl}{top_row+2}:{cr}{top_row+2}",
               note, bg=C["light_grey"], fg=C["dark_grey"],
               bold=False, size=8, h="center", italic=True)


# ---------------------------------------------------------------------------
# Section header bar
# ---------------------------------------------------------------------------
def section_header(ws, rng, title, row_height=20):
    row = int(rng.split(":")[0][1:])
    ws.row_dimensions[row].height = row_height
    merge_cell(ws, rng, title, bg=C["navy"], fg=C["white"],
               bold=True, size=10, h="left")


# ---------------------------------------------------------------------------
# Main dashboard builder
# ---------------------------------------------------------------------------
def build_dashboard(df_p, df_s):

    # ── Calculations ────────────────────────────────────────────────────────
    rev_model = float(df_s["Revenue_EUR"].clip(lower=0).sum())
    rev_ops   = rev_model * (1 - OPS_DISC)

    cf_o  = build_cf(rev_ops)
    cf_m  = build_cf(rev_model)
    cfs_o = [-CAPEX] + [r["fcf"] for r in cf_o]
    cfs_m = [-CAPEX] + [r["fcf"] for r in cf_m]

    irr_o = xirr(cfs_o)
    npv_o = npv(WACC, cfs_o)
    pay_o = payback(cfs_o)

    rev_mw_mo_o = rev_ops / BESS_MW / 12

    mkt_grp = (df_s.groupby("Market_Selected")["Revenue_EUR"]
               .agg(Hours="count", Revenue="sum")
               .sort_values("Revenue", ascending=False))
    mkt_grp["Pct_Rev"] = mkt_grp["Revenue"] / mkt_grp["Revenue"].sum()
    mkt_grp["Pct_Hrs"] = mkt_grp["Hours"]   / mkt_grp["Hours"].sum()

    monthly_soc  = (df_s["Revenue_EUR"].clip(lower=0).resample("ME").sum() / 1e6)
    monthly_fcr  = (df_p["Revenue_FCR_EUR"].resample("ME").sum() / 1e6)
    monthly_spot = (df_p["Revenue_Spot_EUR"].resample("ME").sum() / 1e6)
    spot_stats   = df_p["Price_Spot"].dropna()

    dom_mkt = (df_s.groupby([df_s.index.to_period("M"), "Market_Selected"])
               ["Revenue_EUR"].sum().unstack(fill_value=0).idxmax(axis=1))

    FRIENDLY = {
        "FCR_N":      "FCR-N  (Frequency Reserve – Normal)",
        "FCR_D_UP":   "FCR-D Up  (Frequency Reserve – Disturbance)",
        "FCR_D_DOWN": "FCR-D Down  (Frequency Reserve – Absorption)",
        "mFRR_UP":    "mFRR Up  (Manual Restoration – Up)",
        "mFRR_DOWN":  "mFRR Down  (Manual Restoration – Down)",
        "SPOT_DISC":  "Spot Arbitrage  (Day-ahead discharge)",
        "CHARGING":   "Grid Charging  (Off-peak recharge)",
        "IDLE":       "Idle",
    }
    MKT_BG = {
        "FCR_N":      C["light_blue"],  "FCR_D_UP":   C["light_blue"],
        "FCR_D_DOWN": C["light_blue"],  "mFRR_UP":    C["light_green"],
        "mFRR_DOWN":  C["light_green"], "SPOT_DISC":  C["light_amber"],
        "CHARGING":   C["light_red"],   "IDLE":       C["light_grey"],
    }

    # ── Workbook setup ───────────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws.sheet_view.showGridLines     = False
    ws.sheet_view.showRowColHeaders = False

    # A = left margin | B–H = left panel | I = gutter | J–P = right panel | Q = right margin
    set_col_widths(ws, {
        "A": 1.5,
        "B": 17, "C": 17, "D": 17, "E": 17, "F": 17, "G": 17, "H": 17,
        "I": 2.5,
        "J": 17, "K": 17, "L": 17, "M": 17, "N": 17, "O": 17, "P": 17,
        "Q": 1.5,
    })

    # ── HEADER BANNER ────────────────────────────────────────────────────────
    merge_cell(ws, "B1:P1", bg=C["navy"])
    ws.row_dimensions[1].height = 5

    merge_cell(ws, "B2:P3",
               "PAISTINKULMA BESS  —  PERFORMANCE & INVESTMENT SUMMARY",
               bg=C["navy"], fg=C["white"], bold=True, size=16, h="center")
    ws.row_dimensions[2].height = 28
    ws.row_dimensions[3].height = 28

    merge_cell(ws, "B4:P4",
               "30 MW / 36 MWh Battery Energy Storage  ·  Lempäälä, Finland  ·  "
               "Analysis period: Full Year 2025  ·  "
               "Data: ENTSO-E (spot prices) + Fingrid Open Data (reserve markets)",
               bg=C["mid_blue"], fg=C["white"], bold=False, size=9,
               italic=True, h="center")
    ws.row_dimensions[4].height = 20

    merge_cell(ws, "B5:P5", bg=C["navy"])
    ws.row_dimensions[5].height = 5

    # ── SECTION A: KPI CARDS (rows 7–11) ─────────────────────────────────────
    spacer_row(ws, 6, 16)

    section_header(ws, "B7:H7",  "  KEY PERFORMANCE INDICATORS", row_height=20)
    section_header(ws, "J7:P7",  "  INVESTMENT METRICS  (Operational Estimate, −20% Haircut)",
                   row_height=20)

    ws.row_dimensions[8].height  = 16   # label row
    ws.row_dimensions[9].height  = 32   # value row — big numbers need height
    ws.row_dimensions[10].height = 15   # note row

    kpi_card(ws, 8, 2,
             "Model Revenue  Y1",
             f"€{rev_model/1e6:.2f} M",
             "Perfect foresight upper bound",
             C["mid_blue"], C["mid_blue"], value_size=16)

    kpi_card(ws, 8, 4,
             "Operational Revenue  Y1",
             f"€{rev_ops/1e6:.2f} M",
             "After −20% operational haircut",
             C["teal"], C["teal"], value_size=16)

    kpi_card(ws, 8, 6,
             "Revenue / MW / Month",
             f"€{rev_mw_mo_o:,.0f}",
             "Operational estimate",
             C["mid_blue"], C["mid_blue"], value_size=16)

    kpi_card(ws, 8, 10,
             "Project IRR",
             f"{irr_o:.1%}" if not math.isnan(irr_o) else "N/A",
             "Unlevered, 20-year horizon",
             C["green"], C["green"], value_size=16)

    kpi_card(ws, 8, 12,
             "Project NPV",
             f"€{npv_o/1e6:.1f} M",
             "At 8% WACC",
             C["green"], C["green"], value_size=16)

    kpi_card(ws, 8, 14,
             "Payback Period",
             f"{pay_o:.1f} yrs" if pay_o != float("inf") else "N/A",
             "Undiscounted",
             C["mid_blue"], C["mid_blue"], value_size=16)

    # ── SECTION B: MARKET BREAKDOWN + MONTHLY REVENUE (rows 13+) ─────────────
    spacer_row(ws, 11, 18)
    spacer_row(ws, 12, 4)   # thin accent line above section headers

    ROW_SEC2 = 13
    section_header(ws, f"B{ROW_SEC2}:H{ROW_SEC2}",
                   "  MARKET PARTICIPATION BREAKDOWN", row_height=20)
    section_header(ws, f"J{ROW_SEC2}:P{ROW_SEC2}",
                   "  MONTHLY REVENUE  (€ Millions, SoC-Optimised Strategy)",
                   row_height=20)

    # Market table column headers
    MKT_COLS = ["B", "C", "D", "E", "F", "G"]
    mkt_hdrs = ["Market", "Hours", "% of Year", "Revenue (€)", "% of Rev.", "Avg €/hr"]
    hdr_row = ROW_SEC2 + 1
    ws.row_dimensions[hdr_row].height = 17
    for col_l, hdr in zip(MKT_COLS, mkt_hdrs):
        c = ws[f"{col_l}{hdr_row}"]
        c.value = hdr
        c.fill  = fill(C["mid_blue"])
        c.font  = font(bold=True, color=C["white"], size=9)
        c.alignment = align(h="left" if col_l == "B" else "center")
        c.border = border_thin()

    data_start = hdr_row + 1
    for r_off, (mkt, row_d) in enumerate(mkt_grp.iterrows()):
        er = data_start + r_off
        ws.row_dimensions[er].height = 19
        bg = MKT_BG.get(mkt, C["light_grey"])
        vals = [
            FRIENDLY.get(mkt, mkt),
            int(row_d["Hours"]),
            row_d["Pct_Hrs"],
            row_d["Revenue"],
            row_d["Pct_Rev"],
            row_d["Revenue"] / row_d["Hours"] if row_d["Hours"] else 0,
        ]
        fmts = [None, "#,##0", "0%", "€#,##0", "0%", "€#,##0"]
        for col_l, v, fmt_ in zip(MKT_COLS, vals, fmts):
            c = ws[f"{col_l}{er}"]
            c.value = v
            c.fill  = fill(bg)
            c.font  = font(size=9, bold=(col_l == "B"),
                           color=C["navy"] if col_l == "B" else "000000")
            c.alignment = align(h="left" if col_l == "B" else "center")
            c.border = border_thin()
            if fmt_:
                c.number_format = fmt_

    tot_r = data_start + len(mkt_grp)
    ws.row_dimensions[tot_r].height = 19
    tot_rev = float(mkt_grp["Revenue"].sum())
    tot_hrs = int(mkt_grp["Hours"].sum())
    for col_l, v, fmt_ in zip(
        MKT_COLS,
        ["TOTAL", tot_hrs, 1.0, tot_rev, 1.0, tot_rev / tot_hrs],
        [None, "#,##0", "0%", "€#,##0", "0%", "€#,##0"]
    ):
        c = ws[f"{col_l}{tot_r}"]
        c.value = v
        c.fill  = fill(C["navy"])
        c.font  = font(bold=True, color=C["white"], size=9)
        c.alignment = align(h="left" if col_l == "B" else "center")
        c.border = border_thin()
        if fmt_:
            c.number_format = fmt_

    # Monthly revenue table headers
    MO_COLS  = ["J", "K", "L", "M", "N"]
    mo_hdrs  = ["Month", "Revenue (€ M)", "vs. Spot Naive", "vs. FCR-D Naive", "Dominant Market"]
    ws.row_dimensions[hdr_row].height = 17
    for col_l, hdr in zip(MO_COLS, mo_hdrs):
        c = ws[f"{col_l}{hdr_row}"]
        c.value = hdr
        c.fill  = fill(C["mid_blue"])
        c.font  = font(bold=True, color=C["white"], size=9)
        c.alignment = align(h="left" if col_l == "J" else "center")
        c.border = border_thin()

    MONTH_BG = [C["light_blue"], C["light_grey"]]
    for m_off, (month, mo_val) in enumerate(monthly_soc.items()):
        er = data_start + m_off
        ws.row_dimensions[er].height = 19
        bg      = MONTH_BG[m_off % 2]
        period  = month.to_period("M")
        dom_key = dom_mkt.get(period, "")
        dom_nm  = FRIENDLY.get(dom_key, dom_key).split("(")[0].strip()
        fcr_v   = monthly_fcr.iloc[m_off]  if m_off < len(monthly_fcr)  else 0
        spot_v  = monthly_spot.iloc[m_off] if m_off < len(monthly_spot) else 0
        vs_spot = mo_val - spot_v
        vs_fcr  = mo_val - fcr_v

        row_data = [month.strftime("%B"), mo_val, vs_spot, vs_fcr, dom_nm]
        fmts_    = [None, "0.00", "+0.00;-0.00", "+0.00;-0.00", None]
        for col_l, v, fmt_ in zip(MO_COLS, row_data, fmts_):
            c = ws[f"{col_l}{er}"]
            c.value = v
            c.fill  = fill(bg)
            c.font  = font(
                size=9, bold=(col_l == "J"),
                color=C["navy"]  if col_l == "J"
                      else C["green"] if isinstance(v, float) and v > 0 and col_l in ("L","M")
                      else C["red"]   if isinstance(v, float) and v < 0 and col_l in ("L","M")
                      else "000000")
            c.alignment = align(h="left" if col_l in ("J","N") else "center")
            c.border = border_thin()
            if fmt_:
                c.number_format = fmt_

    mo_tot_r = data_start + len(monthly_soc)
    ws.row_dimensions[mo_tot_r].height = 19
    for col_l, v, fmt_ in zip(
        MO_COLS,
        ["FULL YEAR", monthly_soc.sum(),
         monthly_soc.sum() - monthly_spot.sum(),
         monthly_soc.sum() - monthly_fcr.sum(), ""],
        [None, "0.00", "+0.00;-0.00", "+0.00;-0.00", None]
    ):
        c = ws[f"{col_l}{mo_tot_r}"]
        c.value = v
        c.fill  = fill(C["navy"])
        c.font  = font(bold=True, color=C["white"], size=9)
        c.alignment = align(h="left" if col_l == "J" else "center")
        c.border = border_thin()
        if fmt_:
            c.number_format = fmt_

    # ── SECTION C: ASSUMPTIONS + 20-YEAR CASH FLOW ───────────────────────────
    sec3_start = max(tot_r, mo_tot_r) + 4
    spacer_row(ws, sec3_start - 2, 16)
    spacer_row(ws, sec3_start - 1, 4)

    section_header(ws, f"B{sec3_start}:H{sec3_start}",
                   "  INVESTMENT MODEL ASSUMPTIONS", row_height=20)
    section_header(ws, f"J{sec3_start}:P{sec3_start}",
                   "  20-YEAR FREE CASH FLOW  (€ Millions, Operational Estimate)",
                   row_height=20)

    # Assumption sub-headers
    asmp_hdr = sec3_start + 1
    ws.row_dimensions[asmp_hdr].height = 17
    for col_l, hdr, end_l in zip(["B","E","G"], ["Parameter","Value","Note"],
                                  ["D","F","H"]):
        ws.merge_cells(f"{col_l}{asmp_hdr}:{end_l}{asmp_hdr}")
        c = ws[f"{col_l}{asmp_hdr}"]
        c.value = hdr
        c.fill  = fill(C["mid_blue"])
        c.font  = font(bold=True, color=C["white"], size=9)
        c.alignment = align(h="left" if col_l == "B" else "center")
        c.border = border_thin()

    assumptions = [
        ("CAPEX",              f"€{CAPEX/1e6:.0f} M",     "€400/kW total installed cost"),
        ("OPEX (fixed)",       "€180,000 / yr",            "1.5% of CAPEX annually"),
        ("Revenue Y1",         f"€{rev_ops/1e6:.2f} M",   "Operational estimate (−20% haircut)"),
        ("Revenue growth",     "3.0% p.a.",                "Reserve demand +134% forecast (5 yr)"),
        ("Degradation",        "2.0% p.a.",                "LFP cell capacity loss"),
        ("Cell replacement",   "€3.5 M at year 12",        "Cells only; inverters retained"),
        ("Discount rate",      "8.0% WACC",                "Infrastructure equity target"),
        ("Corporation tax",    "20%",                       "Finland"),
        ("Depreciation",       "10 yr straight-line",      "Initial CAPEX basis"),
        ("Project life",       "20 years",                 ""),
    ]

    for r_off, (param, val, note) in enumerate(assumptions):
        er = asmp_hdr + 1 + r_off
        ws.row_dimensions[er].height = 18
        bg = C["light_grey"] if r_off % 2 == 0 else C["white"]
        for col_l, v, end_l in zip(["B","E","G"], [param, val, note],
                                    ["D","F","H"]):
            ws.merge_cells(f"{col_l}{er}:{end_l}{er}")
            c = ws[f"{col_l}{er}"]
            c.value = v
            c.fill  = fill(bg)
            c.font  = font(size=9, bold=(col_l == "B"))
            c.alignment = align(h="left", v="center")
            c.border = border_thin()

    # CF sub-headers
    cf_hdr = sec3_start + 1
    ws.row_dimensions[cf_hdr].height = 17
    for col_l, hdr in zip(["J","K","L","M","N"],
                           ["Year","Revenue","Free CF","Cumul. CF","Status"]):
        c = ws[f"{col_l}{cf_hdr}"]
        c.value = hdr
        c.fill  = fill(C["mid_blue"])
        c.font  = font(bold=True, color=C["white"], size=9)
        c.alignment = align(h="center")
        c.border = border_thin()

    for r_off, row_d in enumerate(cf_o):
        er      = cf_hdr + 1 + r_off
        cumfcf  = row_d["cumfcf"]
        is_repl = (row_d["yr"] == 12)
        ws.row_dimensions[er].height = 17
        bg = (C["light_amber"] if is_repl
              else C["light_red"]   if cumfcf < 0
              else C["light_grey"]  if r_off % 2 == 0
              else C["white"])
        status = "Cell replacement" if is_repl else ("Payback period" if cumfcf < 0 else "✓")
        vals   = [row_d["yr"], row_d["rev"]/1e6, row_d["fcf"]/1e6, cumfcf/1e6, status]
        fmts_  = ["0", "€0.00", "€0.00", "€0.00", None]
        for col_l, v, fmt_ in zip(["J","K","L","M","N"], vals, fmts_):
            c = ws[f"{col_l}{er}"]
            c.value = v
            c.fill  = fill(bg)
            c.font  = font(
                size=9,
                color=C["red"]   if isinstance(v, float) and v < 0 and col_l in ("L","M")
                      else C["green"] if isinstance(v, float) and v > 0 and col_l == "M"
                      else C["amber"] if is_repl and col_l == "N"
                      else C["green"] if status == "✓" and col_l == "N"
                      else "000000")
            c.alignment = align(h="left" if col_l == "N" else "center")
            c.border = border_thin()
            if fmt_:
                c.number_format = fmt_

    # ── FOOTER ────────────────────────────────────────────────────────────────
    last_cf   = cf_hdr + 20
    last_asmp = asmp_hdr + len(assumptions)
    footer_r  = max(last_cf, last_asmp) + 3

    spacer_row(ws, footer_r - 2, 14)
    spacer_row(ws, footer_r - 1, 4)

    ws.row_dimensions[footer_r].height = 32
    merge_cell(ws, f"B{footer_r}:P{footer_r}",
        "Data sources: ENTSO-E Transparency Platform (FI day-ahead spot prices)  ·  "
        "Fingrid Open Data API: FCR-N (317), FCR-D up (318), FCR-D down (283), "
        "mFRR up (244), mFRR down (106)  ·  "
        "CAPEX benchmark: Ember Jan 2026  ·  "
        "IRR benchmark: 3–7% unlevered Western EU merchant BESS, Capstone DC Nov 2025  ·  "
        "Revenue benchmark: Exilion FI €40,700/MW/month H2 2023, Capalo AI Oct 2024  ·  "
        "Model note: SoC-optimised revenue assumes perfect price foresight; "
        "operational estimate applies −20% haircut for bidding friction and availability.",
        bg=C["light_grey"], fg=C["dark_grey"],
        bold=False, size=8, italic=True, h="left", wrap=True)

    ws.print_area = f"A1:Q{footer_r}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage   = True
    ws.page_setup.fitToWidth  = 1
    ws.page_setup.fitToHeight = 0

    wb.save(DASHBOARD_FILE)
    print(f"[OK] {DASHBOARD_FILE} saved.")


if __name__ == "__main__":
    print("── Generating BESS Dashboard ────────────────────────────")
    df_p, df_s = load_data()
    build_dashboard(df_p, df_s)